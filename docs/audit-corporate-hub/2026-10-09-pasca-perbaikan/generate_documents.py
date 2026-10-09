from pathlib import Path
import subprocess, json, html, hashlib, zipfile, re, os, tempfile
from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Image, Table, TableStyle, PageBreak, Preformatted
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image as PILImage
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
font_cache = Path(tempfile.gettempdir()) / 'p4i-document-cache'
font_cache.mkdir(exist_ok=True)
os.environ.setdefault('XDG_CACHE_HOME', str(font_cache))
NAME = 'laporan-sistem-terbaru'
DIAGRAMS = {
'diagram-arsitektur': r'''digraph G {
rankdir=TB;
browser [label="Pengunjung / browser\nHTML, CSS dan UI interaktif",fillcolor="#dbeafe"];
server [label="Next.js 16.3.8\nProxy nonce/CSP + App Router\nServer rendering per request",fillcolor="#bfdbfe"];
data [label="Konten source\ndata/*.ts + site-config\nNarasi TSX / registry / flags",fillcolor="#dcfce7"];
assets [label="Aset lokal\npublic/ + Inter berlisensi\nnext/image whitelist",fillcolor="#dcfce7"];
media [label="GET media magang\nGuard consent / approval / enabled\nprivate/media + no-store",fillcolor="#fef3c7"];
external [label="Kanal eksternal\nOJS, e-book, WhatsApp, email\nInstagram / Google Maps iframe",fillcolor="#fef3c7"];
ci [label="Pemeliharaan Git + CI\nInstall, audit, lint, typecheck\nBuild + 23 tes",fillcolor="#e0e7ff"];
browser -> server [label="request / respons"];
data -> server [label="import"];
assets -> browser [label="aset / font lokal"];
server -> media [label="route terkontrol"];
media -> browser [label="foto hanya jika layak"];
browser -> external [label="tautan / iframe"];
ci -> data [label="review perubahan",style=dashed];
{rank=same;data;assets;media;}
note [label="Tanpa database / login / RBAC\nGA nonaktif; PR belum merged",shape=note,fillcolor="#f1f5f9"];
server -> note [style=dashed,arrowhead=none];
}''',
'diagram-workflow': r'''digraph G {
rankdir=TB;
start [label="Buka portal",fillcolor="#dbeafe"];
read [label="Baca profil / program / kegiatan\nJejaring / sejarah / magang",fillcolor="#dbeafe"];
choice [label="Tujuan pengunjung",shape=diamond,fillcolor="#fef3c7"];
ui [label="Filter ahli / lightbox\nState sementara, fokus keyboard",fillcolor="#dcfce7"];
contact [label="WhatsApp / email / Gmail\nPesan pada layanan eksternal",fillcolor="#fef3c7"];
publish [label="OJS / e-book\nProses pada platform eksternal",fillcolor="#fef3c7"];
start -> read -> choice;
choice -> ui; choice -> contact; choice -> publish;
subgraph cluster_git {label="Alur pemeliharaan yang tersedia, di luar UI";color="#94a3b8";
curate [label="Kurasi materi / sumber / izin",fillcolor="#e0e7ff"];
edit [label="Edit source dan aset\nCommit / PR / review",fillcolor="#e0e7ff"];
check [label="CI: audit + asset + lint\nTypecheck + build + tests",fillcolor="#e0e7ff"];
release [label="Merge / rilis sesuai izin\nBelum dilakukan untuk PR #1",fillcolor="#e0e7ff"];
curate -> edit -> check -> release;
}
}''',
'diagram-media': r'''digraph G {
rankdir=TB;
request [label="Halaman / preview atau\nGET filename foto magang",fillcolor="#dbeafe"];
flag [label="Fitur internship aktif?",shape=diamond,fillcolor="#fef3c7"];
approval [label="Consent verified AND\napproved for website?",shape=diamond,fillcolor="#fef3c7"];
enabled [label="Item enabled dan\nfilename ada dalam daftar?",shape=diamond,fillcolor="#fef3c7"];
serve [label="Render unoptimized / baca private\nKirim WebP, private no-store",fillcolor="#dcfce7"];
deny [label="Tidak tampil / HTTP 404\nOptimizer path magang: 400",fillcolor="#fee2e2"];
request -> flag;
flag -> approval [label="ya"];
flag -> deny [label="tidak"];
approval -> enabled [label="ya"];
approval -> deny [label="tidak"];
enabled -> serve [label="ya"];
enabled -> deny [label="tidak"];
state [label="Snapshot: consent false\n7 item enabled, 0 foto publik\nStatus izin tidak dipalsukan",shape=note,fillcolor="#f1f5f9"];
state -> approval [style=dashed,arrowhead=none];
}''',
'diagram-erd-institusi': r'''digraph G {
rankdir=TB;
site [label="SITE_CONFIG\nname, legalName, URL\nestablished, social, legal",fillcolor="#dbeafe"];
location [label="LOCATION\nid, label, address, mapQuery\n2 object lokasi",fillcolor="#dcfce7"];
contact [label="CONTACT\nemail, WhatsApp, URL\n1 object kontak",fillcolor="#dcfce7"];
org [label="ORGANIZATION_STRUCTURE\nperiod, historical, source\nData periode 2014",fillcolor="#dbeafe"];
member [label="ORG_MEMBER\nname, title, portrait\nportraitVerified",fillcolor="#dcfce7"];
division [label="BIDANG\nname, members[]\n4 bidang",fillcolor="#dcfce7"];
category [label="PROGRAM_CATEGORY\nid, title, description, icon\n4 kategori",fillcolor="#dbeafe"];
program [label="PROGRAM_NAME\nstring pada programs[]\n23 nama program",fillcolor="#dcfce7"];
site -> location [label="1:N nested"];
site -> contact [label="1:1 nested"];
org -> member [label="1:N kelompok jabatan"];
org -> division [label="1:N nested"];
division -> member [label="1:N members"];
category -> program [label="1:N string[]"];
{rank=same;site;org;category;}
note [label="Object/array TypeScript\nIdentifier konten, tanpa PK/FK SQL",shape=note,fillcolor="#f1f5f9"];
}''',
'diagram-erd-media': r'''digraph G {
rankdir=TB;
intern [label="INTERNSHIP_PROGRAM\nyear, count, approval flags\nConsent saat ini false",fillcolor="#dbeafe"];
photo [label="GALLERY_ITEM\nimage, caption, source, enabled\n7 item media magang",fillcolor="#dcfce7"];
activity [label="ACTIVITY\nid, slug, year, image, CTA\n9 aktivitas",fillcolor="#dbeafe"];
media [label="HISTORICAL_MEDIA\nid, image, source, verification\n11 arsip",fillcolor="#dcfce7"];
history [label="HISTORICAL_MILESTONE\nyear, category, themes\n6 item rinci",fillcolor="#dbeafe"];
expert [label="EXPERT\nid, category, historical, source\n21 profil, koleksi mandiri",fillcolor="#fef3c7"];
milestone [label="MILESTONE\nyear, title, description\n5 item ringkasan mandiri",fillcolor="#fef3c7"];
intern -> photo [label="1:N nested"];
activity -> photo [label="URL opsional + filter",style=dashed];
activity -> media [label="referensi URL gambar",style=dashed];
history -> media [label="0..1 lookup gambar tertentu",style=dashed];
{rank=same;intern;activity;history;}
{rank=same;photo;media;}
{rank=same;expert;milestone;}
note [label="Garis putus-putus: asosiasi kode\nBukan join / foreign key database",shape=note,fillcolor="#f1f5f9"];
}''',
'diagram-erd-registry': r'''digraph G {
rankdir=TB;
feature [label="FEATURE_FLAGS\nfeatureName -> boolean\n8 true, consulting false",fillcolor="#dbeafe"];
route [label="ROUTE\npath, label, optional feature\n11 konten, 10 aktif",fillcolor="#dcfce7"];
platform [label="DIGITAL_PLATFORM\nid, URL, status, feature\nWebsite, OJS, e-book",fillcolor="#dcfce7"];
nav [label="NAV_ITEM\nlabel, href, external, feature",fillcolor="#fef3c7"];
updated [label="CONTENT_UPDATED_AT\nroute -> tanggal editorial\nPartial record",fillcolor="#fef3c7"];
journal [label="JOURNAL\nid, name, URL, active\nISTE aktif",fillcolor="#fef3c7"];
sitemap [label="SITEMAP OUTPUT\nURL aktif + tanggal diketahui\nKeluaran, tidak disimpan di DB",fillcolor="#f1f5f9"];
feature -> route [label="0..1 lookup key",style=dashed];
feature -> platform [label="visibility lookup",style=dashed];
nav -> route [label="href internal",style=dashed];
nav -> platform [label="href eksternal",style=dashed];
platform -> journal [label="base URL jurnal",style=dashed];
updated -> route [label="0..1 tanggal per path",style=dashed];
route -> sitemap [label="filter + generate"];
updated -> sitemap [label="tanggal diketahui"];
}'''
}
for name, body in DIAGRAMS.items():
    header='digraph G {\ngraph [bgcolor="white",pad="0.2",nodesep="0.28",ranksep="0.4",dpi=150];\nnode [shape=box,style="rounded,filled",fontname="DejaVu Sans",fontsize=14,color="#94a3b8",margin="0.12,0.1"];\nedge [fontname="DejaVu Sans",fontsize=11,color="#64748b"];'
    content=body.replace('digraph G {',header,1)
    file=ROOT/(name+'.dot');file.write_text(content)
    for kind in ['png','svg']:
        subprocess.run(['dot', '-T'+kind, str(file), '-o', str(ROOT/(name+'.'+kind))],check=True)

css='''body{font-family:system-ui,sans-serif;color:#0f172a;line-height:1.65;max-width:1100px;margin:auto;padding:40px}h1{font-size:1.8em;border-bottom:2px solid #dbeafe;padding-bottom:12px;margin-top:44px}h2{font-size:1.35em;margin-top:28px}table{border-collapse:collapse;width:100%;font-size:.88em;margin:20px 0}td,th{border:1px solid #cbd5e1;padding:10px;vertical-align:top;text-align:left}th{background:#dbeafe}tr:nth-child(even){background:#f8fafc}img{max-width:100%;max-height:900px;display:block;margin:20px auto}pre{background:#f1f5f9;overflow:auto;padding:20px}code{font-size:.9em;overflow-wrap:anywhere}a{color:#1d4ed8}#TOC{background:#f8fafc;padding:24px;border:1px solid #cbd5e1}@media print{body{padding:0}h1{break-before:page}}'''
(ROOT/'report.css').write_text(css+'\n')
for kind in ['html','docx']:
    options=['--standalone','--toc','--toc-depth=2','--metadata=toc-title:Daftar isi','--data-dir=pandoc-data']
    if kind=='html':options+=['--embed-resources','--css=report.css']
    subprocess.run(['pandoc',NAME+'.md',*options,'-o',NAME+'.'+kind],cwd=ROOT,check=True)

word=Document(ROOT/(NAME+'.docx'))
for section in word.sections:
    section.page_height=Cm(29.7);section.page_width=Cm(21)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
normal=word.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
for heading in ['Heading 1','Heading 2','Heading 3']:
    style = next((item for item in word.styles if item.style_id.lower() == heading.replace(' ', '').lower()), None)
    if style is None: continue
    style.font.color.rgb=RGBColor.from_string('16345A')
    style.paragraph_format.keep_with_next=True
    if heading=='Heading 1': style.paragraph_format.page_break_before=True
for table in word.tables:
    table.autofit=True
    if table.rows:
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(8)
word.core_properties.title='Laporan Sistem Terbaru P4I Corporate Hub'
word.core_properties.subject='Snapshot f77e0f8 setelah perbaikan A01–A16; belum merged'
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');word.settings.element.append(update)
for section in word.sections:
    footer=section.footer.paragraphs[0]
    footer.text='P4I Corporate Hub · 9 Oktober 2026 · Snapshot f77e0f8 · '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8)
word.save(ROOT/(NAME+'.docx'))

fontroot=Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('DV','DejaVuSans.ttf'),('DV-Bold','DejaVuSans-Bold.ttf'),('DV-Italic','DejaVuSans-Oblique.ttf'),('DV-BoldItalic','DejaVuSans-BoldOblique.ttf'),('DVM','DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name,str(fontroot/file)))
pdfmetrics.registerFontFamily('DV',normal='DV',bold='DV-Bold',italic='DV-Italic',boldItalic='DV-BoldItalic')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='Text',fontName='DV',fontSize=9,leading=14,spaceAfter=7))
styles.add(ParagraphStyle(name='SmallText',parent=styles['Text'],fontSize=7.3,leading=10.5,spaceAfter=3,wordWrap='CJK'))
styles.add(ParagraphStyle(name='CodeText',fontName='DVM',fontSize=7,leading=10,spaceAfter=8))
styles.add(ParagraphStyle(name='Cover',fontName='DV-Bold',fontSize=26,leading=34,spaceAfter=20,textColor=colors.HexColor('#16345A')))
for key,size in [('Heading1',17),('Heading2',12),('Heading3',10)]:
    styles[key].fontName='DV-Bold';styles[key].fontSize=size;styles[key].leading=size+5
    styles[key].spaceBefore=12;styles[key].spaceAfter=8;styles[key].keepWithNext=True

class Report(BaseDocTemplate):
    def afterFlowable(self,flowable):
        if isinstance(flowable,Paragraph) and hasattr(flowable,'bookmark_key'):
            key=flowable.bookmark_key;level=flowable.heading_level
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(),key,level=level,closed=True)
            self.notify('TOCEntry',(level,flowable.getPlainText(),self.page,key))

def inline(items):
    result=[]
    for item in items:
        t,c=item['t'],item.get('c')
        if t=='Str':result.append(html.escape(c))
        elif t in ['Space','SoftBreak']:result.append(' ')
        elif t=='LineBreak':result.append('<br/>')
        elif t=='Strong':result.append('<b>'+inline(c)+'</b>')
        elif t=='Emph':result.append('<i>'+inline(c)+'</i>')
        elif t=='Code':result.append('<font name="DVM">'+html.escape(c[1])+'</font>')
        elif t=='Link':
            target=c[2][0]
            if target.startswith(('https://','http://')):result.append('<link href="'+html.escape(target,quote=True)+'" color="#1d4ed8">'+inline(c[1])+'</link>')
            else:result.append(inline(c[1]))
        elif t=='Image':result.append(inline(c[1]))
        elif t=='Quoted':result.append('“'+inline(c[1])+'”')
    return ''.join(result)

width=A4[0]-84

def render(blocks):
    result=[]
    for block in blocks:
        t,c=block['t'],block.get('c')
        if t=='Header':
            level=min(c[0],3)
            if level==1:result.append(PageBreak())
            p=Paragraph(inline(c[2]),styles['Heading'+str(level)])
            if level<=2:p.bookmark_key=c[1][0];p.heading_level=level-1
            result.append(p)
        elif t in ['Para','Plain']:
            if len(c)==1 and c[0]['t']=='Image':
                file=ROOT/c[0]['c'][2][0]
                with PILImage.open(file) as im:iw,ih=im.size
                factor=min(width/iw,555/ih)
                result.extend([Image(str(file),width=iw*factor,height=ih*factor),Paragraph(inline(c[0]['c'][1]),styles['SmallText']),Spacer(1,8)])
            else:result.append(Paragraph(inline(c),styles['Text']))
        elif t=='Figure':result.extend(render(c[2]))
        elif t in ['BulletList','OrderedList']:
            items=c if t=='BulletList' else c[1]
            for idx,item in enumerate(items,1):
                for j,sub in enumerate(item):
                    if sub['t'] in ['Para','Plain']:
                        prefix=('• ' if t=='BulletList' else str(idx)+'. ') if j==0 else ''
                        result.append(Paragraph(prefix+inline(sub['c']),styles['Text']))
                    else:result.extend(render([sub]))
        elif t=='CodeBlock':result.append(Preformatted(c[1],styles['CodeText'],maxLineLength=100))
        elif t=='Table':
            allrows=list(c[3][1])
            for body in c[4]:allrows.extend(body[2]+body[3])
            allrows.extend(c[5][1]);rows=[]
            for row in allrows:
                cells=[]
                for cell in row[1]:
                    parts=[inline(x['c']) for x in cell[4] if x['t'] in ['Plain','Para']]
                    cells.append(Paragraph('<br/>'.join(parts),styles['SmallText']))
                rows.append(cells)
            if rows:
                n=len(rows[0]);weights=[.26,.74] if n==2 else [.23,.37,.40] if n==3 else [1/n]*n
                table=Table(rows,colWidths=[width*x for x in weights],repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dbeafe')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#cbd5e1')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')])]))
                result.extend([table,Spacer(1,10)])
        elif t=='BlockQuote':result.extend(render(c))
    return result

ast=json.loads(subprocess.check_output(['pandoc',NAME+'.md','-t','json'],cwd=ROOT))
story=[Spacer(1,50),Paragraph('Laporan Analisis Sistem<br/>dan Audit Terbaru',styles['Cover']),Paragraph('P4I Corporate Hub',styles['Cover']),Paragraph('Kondisi implementasi setelah perbaikan A01–A16',styles['Heading2']),Spacer(1,20),Paragraph('9 Oktober 2026 · UTC<br/>Repository HannR002/p4i-corporate-hub<br/>Snapshot f77e0f8 · fix/corporate-hub-audit',styles['Text']),Spacer(1,22),Paragraph('Workflow · ERD aktual · Aktor dan use case · Fitur<br/>Struktur kode · Teknologi · Keamanan · Hasil pengujian',styles['Text']),Spacer(1,25),Paragraph('<b>Status:</b> PR #1 masih draft; belum digabungkan ke main.<br/>Laporan berbasis kode, log CI, dan audit dependency.<br/>Tidak memuat tahap rekomendasi pengembangan baru.',styles['Text']),PageBreak(),Paragraph('Daftar isi',styles['Heading1'])]
toc=TableOfContents();toc.levelStyles=[ParagraphStyle(name='TOC1',fontName='DV',fontSize=9,leading=13,spaceBefore=7),ParagraphStyle(name='TOC2',fontName='DV',fontSize=8,leading=11,leftIndent=12,spaceBefore=3)]
story.append(toc);story.extend(render(ast['blocks']))

def footer(canvas,doc):
    canvas.saveState();canvas.setFont('DV',7);canvas.setFillColor(colors.HexColor('#64748b'))
    canvas.drawString(42,25,'P4I Corporate Hub · Snapshot f77e0f8 · 9 Oktober 2026')
    canvas.drawRightString(A4[0]-42,25,str(doc.page));canvas.restoreState()
doc=Report(str(ROOT/(NAME+'.pdf')),pagesize=A4,leftMargin=42,rightMargin=42,topMargin=42,bottomMargin=42,title='Laporan Sistem Terbaru P4I Corporate Hub',author='Analisis berbasis repository')
doc.addPageTemplates(PageTemplate(id='normal',frames=[Frame(42,42,width,A4[1]-84,id='content',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=footer))
doc.multiBuild(story)
(ROOT/'requirements-dokumen.txt').write_text('reportlab\npython-docx\nPillow\n\n# Tool sistem tambahan: pandoc dan graphviz (dot)\n')
manifest={str(file.relative_to(ROOT)):hashlib.sha256(file.read_bytes()).hexdigest() for file in sorted(ROOT.rglob('*')) if file.is_file() and file.name not in ['manifest-sha256.json','paket-laporan-terbaru.zip']}
(ROOT/'manifest-sha256.json').write_text(json.dumps(manifest,indent=2)+'\n')
with zipfile.ZipFile(ROOT/'paket-laporan-terbaru.zip','w',zipfile.ZIP_DEFLATED) as archive:
    for file in sorted(ROOT.rglob('*')):
        if file.is_file() and file.name!='paket-laporan-terbaru.zip':archive.write(file,str(file.relative_to(ROOT)))
print(json.dumps({'generated':['PDF','DOCX','HTML','Markdown'],'diagrams':len(DIAGRAMS),'bundle_bytes':(ROOT/'paket-laporan-terbaru.zip').stat().st_size},indent=2))
